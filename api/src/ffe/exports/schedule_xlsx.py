"""Schedule workbook (ADR-0009): one row per placement grouped by zone, « Budget » sheet for the
internal variant. Every amount is a value from the summary (money.py); no Excel formula."""

from datetime import date
from io import BytesIO
from typing import Any, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy import select
from sqlalchemy.orm import Session

from ffe.i18n import fr
from ffe.models import Project, ProjectItem, Supplier, Zone
from ffe.services.budget import summary_with_budget
from ffe.services.project_item import list_items

Variant = Literal["internal", "client"]
MONEY_FORMAT = "# ##0.00"
QTY_FORMAT = "# ##0.###"

INTERNAL_COLUMNS: tuple[str, ...] = (
    "Zone",
    "Nom de zone",
    "×",
    "Code",
    "Désignation",
    "Type",
    "Finition",
    "Fournisseur",
    "Qté / zone",
    "Qté totale",
    "PU net (devise)",
    "Devise",
    "PU net HT €",
    "Total HT €",
    "Statut",
    "Délai (sem.)",
    "Commander avant le",
    "Notes internes",
)
CLIENT_COLUMNS: tuple[str, ...] = (
    "Zone",
    "Nom de zone",
    "×",
    "Code",
    "Désignation",
    "Type",
    "Finition",
    "Qté / zone",
    "Qté totale",
    "Statut",
    "Description client",
)
BUDGET_COLUMNS: tuple[str, ...] = (
    "Périmètre",
    "×",
    "Budget HT €",
    "Prescrit HT €",
    "Validé HT €",
    "Engagé HT €",
    "Reste HT €",
    "Alerte",
)


def cents(value: int | None) -> float | None:
    """Centimes → nombre à deux décimales (valeur, jamais une formule)."""
    return None if value is None else value / 100


def milli(value: int) -> float:
    return value / 1000


def _header(sheet: Worksheet, columns: tuple[str, ...]) -> None:
    sheet.append(list(columns))
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"


def _widths(sheet: Worksheet, widths: dict[int, int]) -> None:
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = width


def build_schedule(session: Session, project_id: int, variant: Variant, today: date) -> bytes:
    project = session.get(Project, project_id)
    assert project is not None
    summary = summary_with_budget(session, project_id, today)
    amounts: dict[int, dict[str, Any]] = {
        line["project_item_id"]: line for line in summary["lines"]
    }
    zones = {z.id: z for z in session.scalars(select(Zone).where(Zone.project_id == project_id))}
    suppliers = {s.id: s.name for s in session.scalars(select(Supplier))}
    items: list[ProjectItem] = [
        i for i in list_items(session, project_id) if i.status not in ("cancelled", "replaced")
    ]
    if variant == "client":
        items = [i for i in items if i.kind != "fee"]

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Schedule"
    columns = INTERNAL_COLUMNS if variant == "internal" else CLIENT_COLUMNS
    _header(sheet, columns)

    rows: list[tuple[tuple[int, int, int], list[Any]]] = []
    for item in items:
        line = amounts[item.id]
        placements = item.placements or []
        for placement in placements:
            zone = zones.get(placement.zone_id) if placement.zone_id is not None else None
            sort_key = (zone.sort_order if zone else 10**9, zone.id if zone else 0, item.id)
            multiplier = zone.multiplier if zone else 1
            common: list[Any] = [
                zone.code if zone else "",
                zone.name if zone else "Non affecté",
                multiplier,
                item.item_code,
                item.name,
                fr.ITEM_KIND_LABELS[item.kind],
                item.finish,
            ]
            if variant == "internal":
                values = common + [
                    suppliers.get(item.supplier_id, "") if item.supplier_id else "",
                    milli(placement.qty_milli),
                    milli(line["total_qty_milli"]),
                    cents(line["net_unit_cents"]) if item.kind != "existing" else None,
                    item.currency,
                    cents(line["net_unit_eur_cents"]) if item.kind != "existing" else None,
                    cents(line["total_ht_cents"]) if item.kind != "existing" else None,
                    fr.ITEM_STATUS_LABELS[item.status],
                    item.lead_time_weeks,
                    line["order_by"],
                    item.internal_notes,
                ]
            else:
                values = common + [
                    milli(placement.qty_milli),
                    milli(line["total_qty_milli"]),
                    fr.ITEM_STATUS_LABELS[item.status],
                    item.client_description,
                ]
            rows.append((sort_key, values))
    for _key, values in sorted(rows, key=lambda r: r[0]):
        sheet.append(values)

    if variant == "internal":
        money_columns = (11, 13, 14)
        qty_columns = (9, 10)
        for row in sheet.iter_rows(min_row=2):
            for index in money_columns:
                row[index - 1].number_format = MONEY_FORMAT
            for index in qty_columns:
                row[index - 1].number_format = QTY_FORMAT
        total_row = len(rows) + 2
        sheet.cell(row=total_row, column=13, value="Prescrit HT €").font = Font(bold=True)
        total = sheet.cell(row=total_row, column=14, value=cents(summary["specified_cents"]))
        total.font = Font(bold=True)
        total.number_format = MONEY_FORMAT
        sheet.cell(row=total_row + 1, column=13, value="TVA €")
        sheet.cell(
            row=total_row + 1, column=14, value=cents(summary["vat_cents"])
        ).number_format = MONEY_FORMAT
        sheet.cell(row=total_row + 2, column=13, value="TTC €")
        sheet.cell(
            row=total_row + 2, column=14, value=cents(summary["ttc_cents"])
        ).number_format = MONEY_FORMAT
        if summary["ose_cents"]:
            sheet.cell(row=total_row + 3, column=13, value="OS&E hors enveloppe HT €")
            sheet.cell(
                row=total_row + 3, column=14, value=cents(summary["ose_cents"])
            ).number_format = MONEY_FORMAT
        _widths(sheet, {1: 8, 2: 22, 4: 10, 5: 36, 6: 10, 7: 14, 8: 18, 18: 40})
        _budget_sheet(workbook.create_sheet("Budget"), summary["budget"])
    else:
        for row in sheet.iter_rows(min_row=2):
            for index in (8, 9):
                row[index - 1].number_format = QTY_FORMAT
        _widths(sheet, {1: 8, 2: 22, 4: 10, 5: 36, 6: 10, 7: 14, 11: 50})

    info = workbook.create_sheet("Projet")
    info.append(["Projet", f"{project.code} — {project.name}"])
    info.append(["Variante", "interne" if variant == "internal" else "client"])
    info.append(["Émis le", today.isoformat()])
    info.append(["Montants", "HT, EUR, net de remise ; calculés par FFE, sans formule"])

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _budget_sheet(sheet: Worksheet, budget: dict[str, Any]) -> None:
    _header(sheet, BUDGET_COLUMNS)

    def row(label: str, multiplier: int, r: dict[str, Any]) -> None:
        sheet.append(
            [
                label,
                multiplier,
                cents(r["budget_cents"]),
                cents(r["specified_cents"]),
                cents(r["approved_cents"]),
                cents(r["engaged_cents"]),
                cents(r["remaining_cents"]),
                {"none": "", "warning": "≥ 90 %", "over": "≥ 100 %"}[r["alert"]],
            ]
        )

    row("Global", 1, {**budget, "budget_cents": budget["global_cents"]})
    for r in budget["by_zone"]:
        row(f"Zone {r['label']}", r["multiplier"], r)
    for r in budget["by_category"]:
        row(f"Catégorie {r['label']}", 1, r)
    for sheet_row in sheet.iter_rows(min_row=2):
        for index in (3, 4, 5, 6, 7):
            sheet_row[index - 1].number_format = MONEY_FORMAT
    _widths(sheet, {1: 32})
